import docx
import pytest

from gerar_manuais import (
    Callout,
    Codigo,
    ErroDeFonte,
    Lista,
    Manual,
    Paragrafo,
    Tabela,
    Titulo,
    gerar_docx,
    main,
    partir_inline,
    parse,
    nome_saida,
    numerar,
    validar,
)

FONTE_MINIMA = """---
titulo: Manual de Teste
subtitulo: Subtítulo do manual
versao: "1.0"
data: 2026-08-18
---

## Introdução

Primeiro parágrafo.

### Uma seção

Segundo parágrafo.
"""


def test_front_matter_vira_meta():
    manual = parse(FONTE_MINIMA)
    assert manual.meta["titulo"] == "Manual de Teste"
    assert manual.meta["subtitulo"] == "Subtítulo do manual"
    assert manual.meta["versao"] == "1.0"
    assert manual.meta["data"] == "2026-08-18"


def test_titulos_viram_blocos_com_nivel():
    blocos = parse(FONTE_MINIMA).blocos
    titulos = [b for b in blocos if isinstance(b, Titulo)]
    assert [(t.nivel, t.texto) for t in titulos] == [
        (1, "Introdução"),
        (2, "Uma seção"),
    ]


def test_paragrafos_preservam_ordem():
    blocos = parse(FONTE_MINIMA).blocos
    paragrafos = [b.texto for b in blocos if isinstance(b, Paragrafo)]
    assert paragrafos == ["Primeiro parágrafo.", "Segundo parágrafo."]


def test_linhas_seguidas_viram_um_paragrafo_so():
    fonte = FONTE_MINIMA.replace("Primeiro parágrafo.", "Uma linha\ne a continuação dela")
    paragrafos = [b.texto for b in parse(fonte).blocos if isinstance(b, Paragrafo)]
    assert paragrafos[0] == "Uma linha e a continuação dela"


def test_fonte_sem_front_matter_falha():
    with pytest.raises(ErroDeFonte, match="front-matter"):
        parse("## Introdução\n\nTexto.\n")


def test_front_matter_incompleto_falha():
    fonte = FONTE_MINIMA.replace('versao: "1.0"\n', "")
    with pytest.raises(ErroDeFonte, match="versao"):
        parse(fonte)


CABECALHO = """---
titulo: T
subtitulo: S
versao: "1.0"
data: 2026-08-18
---

"""


def test_tabela_com_cabecalho_e_linhas():
    fonte = CABECALHO + (
        "| Peça | Onde fica |\n"
        "|---|---|\n"
        "| scpi-api | systemd |\n"
        "| nginx | /etc/nginx |\n"
    )
    tabela = [b for b in parse(fonte).blocos if isinstance(b, Tabela)][0]
    assert tabela.cabecalho == ["Peça", "Onde fica"]
    assert tabela.linhas == [["scpi-api", "systemd"], ["nginx", "/etc/nginx"]]


def test_tabela_com_numero_de_colunas_errado_falha():
    fonte = CABECALHO + "| A | B |\n|---|---|\n| só uma |\n"
    with pytest.raises(ErroDeFonte, match="colunas"):
        parse(fonte)


def test_tabela_com_separador_espacado():
    fonte = CABECALHO + (
        "| Peça | Onde fica |\n"
        "| --- | --- |\n"
        "| scpi-api | systemd |\n"
        "| nginx | /etc/nginx |\n"
    )
    tabela = [b for b in parse(fonte).blocos if isinstance(b, Tabela)][0]
    assert tabela.cabecalho == ["Peça", "Onde fica"]
    assert tabela.linhas == [["scpi-api", "systemd"], ["nginx", "/etc/nginx"]]


def test_tabela_com_separador_alinhamento():
    fonte = CABECALHO + (
        "| Esquerda | Centro | Direita |\n"
        "| :--- | :---: | ---: |\n"
        "| A | B | C |\n"
    )
    tabela = [b for b in parse(fonte).blocos if isinstance(b, Tabela)][0]
    assert tabela.cabecalho == ["Esquerda", "Centro", "Direita"]
    assert tabela.linhas == [["A", "B", "C"]]


def test_linha_com_pipe_mas_separador_invalido_vira_paragrafo():
    fonte = CABECALHO + "| abc | def |\n| xyz | uvw |\n"
    blocos = parse(fonte).blocos
    assert len(blocos) == 1
    assert isinstance(blocos[0], Paragrafo)
    assert "| abc | def |" in blocos[0].texto


def test_bloco_de_codigo_preserva_linhas_e_linguagem():
    fonte = CABECALHO + "```bash\ncd /opt/scpi\ngit pull\n```\n"
    codigo = [b for b in parse(fonte).blocos if isinstance(b, Codigo)][0]
    assert codigo.linguagem == "bash"
    assert codigo.linhas == ["cd /opt/scpi", "git pull"]


def test_codigo_nao_vira_paragrafo():
    fonte = CABECALHO + "```\n## isto nao e titulo\n```\n"
    assert not [b for b in parse(fonte).blocos if isinstance(b, Titulo)]


def test_bloco_de_codigo_sem_fechamento_falha():
    fonte = CABECALHO + "```bash\ncd /opt/scpi\n"
    with pytest.raises(ErroDeFonte, match="não fechado"):
        parse(fonte)


def test_callout_reconhece_rotulo():
    fonte = CABECALHO + "> **Pegadinha:** o plano gratuito usa HEAD.\n"
    callout = [b for b in parse(fonte).blocos if isinstance(b, Callout)][0]
    assert callout.rotulo == "Pegadinha"
    assert callout.texto == "o plano gratuito usa HEAD."


def test_citacao_sem_rotulo_conhecido_vira_callout_sem_rotulo():
    fonte = CABECALHO + "> texto solto de citação.\n"
    callout = [b for b in parse(fonte).blocos if isinstance(b, Callout)][0]
    assert callout.rotulo == ""
    assert callout.texto == "texto solto de citação."


def test_listas_marcada_e_numerada():
    fonte = CABECALHO + "- primeiro\n- segundo\n\n1. um\n2. dois\n"
    listas = [b for b in parse(fonte).blocos if isinstance(b, Lista)]
    assert listas[0].ordenada is False and listas[0].itens == ["primeiro", "segundo"]
    assert listas[1].ordenada is True and listas[1].itens == ["um", "dois"]


def test_partir_inline_separa_negrito_e_codigo():
    assert partir_inline("use **git pull** no `/opt/scpi` agora") == [
        ("normal", "use "),
        ("negrito", "git pull"),
        ("normal", " no "),
        ("codigo", "/opt/scpi"),
        ("normal", " agora"),
    ]


def test_partir_inline_sem_marcacao():
    assert partir_inline("texto simples") == [("normal", "texto simples")]


def test_numeracao_hierarquica():
    blocos = [
        Titulo(1, "Introdução"),
        Titulo(2, "Visão geral"),
        Titulo(1, "Backup"),
        Titulo(2, "Peças"),
        Titulo(3, "Na VM"),
        Titulo(2, "Comandos"),
    ]
    assert [t.numero for t in numerar(blocos)] == ["1", "1.1", "2", "2.1", "2.1.1", "2.2"]


def test_numeracao_ignora_blocos_que_nao_sao_titulo():
    blocos = [Titulo(1, "A"), Paragrafo("texto"), Titulo(1, "B")]
    assert [b.numero for b in numerar(blocos) if isinstance(b, Titulo)] == ["1", "2"]


def test_validar_avisa_sobre_confirmar_pendente():
    manual = Manual(meta={}, blocos=[Paragrafo("O OCID é ⚠️ CONFIRMAR com o Gustavo.")])
    avisos = validar(manual, fontes_conhecidas=set())
    assert any("CONFIRMAR" in a for a in avisos)


def test_validar_aceita_referencia_para_manual_existente():
    manual = Manual(meta={}, blocos=[Paragrafo("Ver 07-banco-de-dados.md, seção Migrations.")])
    assert validar(manual, fontes_conhecidas={"07-banco-de-dados.md"}) == []


def test_validar_falha_em_referencia_para_manual_inexistente():
    manual = Manual(meta={}, blocos=[Paragrafo("Ver 99-inexistente.md para detalhes.")])
    with pytest.raises(ErroDeFonte, match="99-inexistente.md"):
        validar(manual, fontes_conhecidas={"07-banco-de-dados.md"})


def test_nome_de_saida_derivado_da_fonte():
    assert nome_saida("01-ambiente-dev.md") == "SCPI-Manual-Ambiente-Dev.docx"
    assert nome_saida("11-fluxo-ci.md") == "SCPI-Manual-Fluxo-Ci.docx"


def test_numeracao_falha_ao_pular_nivel():
    blocos = [Titulo(1, "A"), Titulo(3, "B")]
    with pytest.raises(ErroDeFonte, match="B"):
        numerar(blocos)


def test_numeracao_falha_ao_comcar_em_nivel_2():
    blocos = [Titulo(2, "X")]
    with pytest.raises(ErroDeFonte):
        numerar(blocos)


def test_gerar_docx_produz_arquivo_legivel(tmp_path):
    fonte = CABECALHO + (
        "## Introdução\n\nTexto de abertura.\n\n"
        "### Peças\n\n"
        "| Peça | Função |\n|---|---|\n| scpi-api | serve a API |\n\n"
        "> **Importante:** não remover o HEAD da rota.\n\n"
        "```bash\nsystemctl status scpi-api\n```\n"
    )
    manual = parse(fonte)
    numerar(manual.blocos)
    destino = tmp_path / "saida.docx"
    gerar_docx(manual, destino)

    assert destino.exists()
    documento = docx.Document(str(destino))
    textos = [p.text for p in documento.paragraphs]
    assert "T" in textos[0]
    assert "1 Introdução" in textos
    assert "1.1 Peças" in textos
    assert any("Importante" in t for t in textos)
    assert any("systemctl status scpi-api" in t for t in textos)
    assert len(documento.tables) == 1
    assert documento.tables[0].cell(0, 0).text == "Peça"


def test_estilos_de_titulo_sao_heading_do_word(tmp_path):
    manual = parse(CABECALHO + "## Um\n\ntexto\n\n### Dois\n\ntexto\n")
    numerar(manual.blocos)
    destino = tmp_path / "s.docx"
    gerar_docx(manual, destino)
    estilos = {p.text: p.style.name for p in docx.Document(str(destino)).paragraphs}
    assert estilos["1 Um"] == "Heading 1"
    assert estilos["1.1 Dois"] == "Heading 2"


def test_tabela_com_pipe_escapado_produz_pipe_literal():
    fonte = CABECALHO + (
        "| Comando | Para quê |\n"
        "|---|---|\n"
        "| `npm audit --json \\| node gate.mjs` | roda o gate |\n"
    )
    tabela = [b for b in parse(fonte).blocos if isinstance(b, Tabela)][0]
    assert tabela.cabecalho == ["Comando", "Para quê"]
    assert len(tabela.linhas) == 1
    assert len(tabela.linhas[0]) == 2
    # A célula deve conter o pipe literal sem a barra invertida
    assert tabela.linhas[0][0] == "`npm audit --json | node gate.mjs`"
    assert tabela.linhas[0][1] == "roda o gate"


def test_tabela_com_pipe_escapado_valida_colunas_erradas():
    fonte = CABECALHO + (
        "| A | B |\n"
        "|---|---|\n"
        "| `só tem uma \\| com escape` |\n"
    )
    with pytest.raises(ErroDeFonte, match="colunas"):
        parse(fonte)


# --- B1: callout de várias linhas vira um único Callout ---------------------


def test_callout_de_varias_linhas_vira_um_callout_so():
    fonte = CABECALHO + (
        "> **Atenção:** primeira linha do aviso que\n"
        "> continua na segunda linha e\n"
        "> termina na terceira.\n"
    )
    callouts = [b for b in parse(fonte).blocos if isinstance(b, Callout)]
    assert len(callouts) == 1
    assert callouts[0].rotulo == "Atenção"
    assert callouts[0].texto == (
        "primeira linha do aviso que continua na segunda linha e termina na terceira."
    )


def test_callout_sem_rotulo_de_varias_linhas_tambem_junta():
    fonte = CABECALHO + (
        "> texto solto que\n"
        "> continua aqui.\n"
    )
    callouts = [b for b in parse(fonte).blocos if isinstance(b, Callout)]
    assert len(callouts) == 1
    assert callouts[0].rotulo == ""
    assert callouts[0].texto == "texto solto que continua aqui."


# --- B2: bloco de código dentro de citação vira Codigo normal ---------------


def test_bloco_de_codigo_dentro_de_citacao_vira_codigo_normal():
    fonte = CABECALHO + (
        "> **Atenção:** confira antes de commitar:\n"
        ">\n"
        "> ```bash\n"
        "> git diff -- a b\n"
        "> ```\n"
        ">\n"
        "> A saída tem de estar vazia.\n"
    )
    blocos = parse(fonte).blocos
    assert [type(b).__name__ for b in blocos] == ["Callout", "Codigo", "Callout"]
    callout1, codigo, callout2 = blocos
    assert callout1.rotulo == "Atenção"
    assert callout1.texto == "confira antes de commitar:"
    assert codigo.linguagem == "bash"
    assert codigo.linhas == ["git diff -- a b"]
    assert callout2.rotulo == ""
    assert callout2.texto == "A saída tem de estar vazia."
    # nenhuma crase de bloco de código deve sobrar no texto do callout
    assert "```" not in callout1.texto
    assert "```" not in callout2.texto


def test_codigo_dentro_de_citacao_sem_fechamento_falha():
    fonte = CABECALHO + (
        "> **Atenção:** confira antes de commitar:\n"
        "> ```bash\n"
        "> git diff -- a b\n"
    )
    with pytest.raises(ErroDeFonte, match="não fechado"):
        parse(fonte)


# --- B3: negrito com código aninhado ----------------------------------------


def test_partir_inline_negrito_com_codigo_aninhado():
    assert partir_inline("**Incrementar o `?v=N`**") == [
        ("negrito", "Incrementar o "),
        ("negrito-codigo", "?v=N"),
    ]


def test_partir_inline_negrito_todo_ele_codigo():
    assert partir_inline("**`BackEnd/infra/migrations.py`**") == [
        ("negrito-codigo", "BackEnd/infra/migrations.py"),
    ]


def test_escrever_inline_negrito_com_codigo_produz_runs_sem_crase(tmp_path):
    fonte = CABECALHO + "## Um\n\ntexto com **`BackEnd/infra/migrations.py`** no meio.\n"
    manual = parse(fonte)
    numerar(manual.blocos)
    destino = tmp_path / "s.docx"
    gerar_docx(manual, destino)
    documento = docx.Document(str(destino))
    paragrafo = [p for p in documento.paragraphs if "no meio" in p.text][0]
    assert "`" not in paragrafo.text
    runs_codigo = [r for r in paragrafo.runs if r.text == "BackEnd/infra/migrations.py"]
    assert runs_codigo, "esperava um run isolado com o texto do código"
    assert runs_codigo[0].bold is True
    assert runs_codigo[0].font.name == "Consolas"


# --- B4: CONFIRMAR dentro de Codigo entra na contagem, mas não na referência cruzada ---


def test_validar_conta_confirmar_dentro_de_bloco_de_codigo():
    manual = Manual(
        meta={}, blocos=[Codigo("bash", ["oci os bucket create --name ⚠️ CONFIRMAR"])]
    )
    avisos = validar(manual, fontes_conhecidas=set())
    assert any("CONFIRMAR" in a for a in avisos)


def test_validar_nao_varre_codigo_para_referencia_cruzada():
    manual = Manual(
        meta={},
        blocos=[Codigo("bash", ["python docs/gerar_manuais.py 99-inexistente.md"])],
    )
    # não deve levantar, mesmo citando um "manual" inexistente dentro do comando
    assert validar(manual, fontes_conhecidas=set()) == []


# --- B5: main() aceita as duas formas de caminho e falha com mensagem clara ---


def test_main_aceita_nome_solto():
    assert main(["01-ambiente-dev.md"]) == 0


def test_main_aceita_caminho_relativo_a_raiz_do_repo():
    assert main(["docs/01-ambiente-dev.md"]) == 0


def test_main_com_arquivo_inexistente_devolve_mensagem_clara(capsys):
    codigo = main(["99-nao-existe.md"])
    assert codigo == 1
    saida = capsys.readouterr()
    assert "Traceback" not in saida.out
    assert "Traceback" not in saida.err
    assert "99-nao-existe.md" in (saida.out + saida.err)


# --- C1: regerar sem mudar a fonte produz bytes idênticos -------------------


def test_gerar_docx_duas_vezes_produz_bytes_identicos(tmp_path, monkeypatch):
    # Sem controlar o relógio, duas gerações rápidas em sequência podem cair
    # no mesmo segundo par e "passar" mesmo com o bug (o zip trunca a
    # resolução do timestamp em 2s). Avançamos o relógio manualmente entre
    # as duas gerações para que o teste pegue o bug de verdade.
    import time as time_mod

    relogio = [1_700_000_000]

    def tempo_falso():
        relogio[0] += 10
        return relogio[0]

    monkeypatch.setattr(time_mod, "time", tempo_falso)

    fonte = CABECALHO + (
        "## Introdução\n\nTexto de abertura.\n\n"
        "### Peças\n\n"
        "| Peça | Função |\n|---|---|\n| scpi-api | serve a API |\n\n"
        "> **Importante:** não remover o HEAD da rota.\n\n"
        "```bash\nsystemctl status scpi-api\n```\n"
    )
    destino1 = tmp_path / "um.docx"
    destino2 = tmp_path / "dois.docx"

    manual1 = parse(fonte)
    numerar(manual1.blocos)
    gerar_docx(manual1, destino1)

    manual2 = parse(fonte)
    numerar(manual2.blocos)
    gerar_docx(manual2, destino2)

    assert destino1.read_bytes() == destino2.read_bytes()
