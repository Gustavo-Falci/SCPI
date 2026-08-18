"""Gera os manuais .docx do SCPI a partir das fontes Markdown em docs/.

Uso:
    python docs/gerar_manuais.py                    # regera todos
    python docs/gerar_manuais.py 01-ambiente-dev.md # regera um
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CAMPOS_OBRIGATORIOS = ("titulo", "subtitulo", "versao", "data")
ROTULOS_CALLOUT = ("Importante", "Pegadinha", "Quando usar", "Atenção")
RAIZ_DOCS = Path(__file__).resolve().parent
DESTINO = RAIZ_DOCS / "dist"
FONTE_CODIGO = "Consolas"
CINZA_FUNDO = "F2F2F2"


class ErroDeFonte(Exception):
    """Fonte Markdown malformada. Aborta a geração."""


@dataclass
class Titulo:
    nivel: int
    texto: str
    numero: str = ""


@dataclass
class Paragrafo:
    texto: str


@dataclass
class Tabela:
    cabecalho: list
    linhas: list


@dataclass
class Codigo:
    linguagem: str
    linhas: list


@dataclass
class Callout:
    rotulo: str
    texto: str


@dataclass
class Lista:
    itens: list
    ordenada: bool


@dataclass
class Manual:
    meta: dict
    blocos: list = field(default_factory=list)


def _parse_front_matter(texto: str) -> tuple[dict, str]:
    if not texto.startswith("---\n"):
        raise ErroDeFonte("fonte sem front-matter: primeira linha deve ser '---'")
    fim = texto.find("\n---\n", 4)
    if fim == -1:
        raise ErroDeFonte("front-matter sem fechamento '---'")
    meta = {}
    for linha in texto[4:fim].splitlines():
        if not linha.strip():
            continue
        chave, _, valor = linha.partition(":")
        meta[chave.strip()] = valor.strip().strip('"')
    faltando = [c for c in CAMPOS_OBRIGATORIOS if not meta.get(c)]
    if faltando:
        raise ErroDeFonte("front-matter incompleto, faltam: %s" % ", ".join(faltando))
    return meta, texto[fim + 5:]


def partir_inline(texto: str) -> list:
    partes = []
    for pedaco in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
        if not pedaco:
            continue
        if pedaco.startswith("**") and pedaco.endswith("**"):
            partes.append(("negrito", pedaco[2:-2]))
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            partes.append(("codigo", pedaco[1:-1]))
        else:
            partes.append(("normal", pedaco))
    return partes


def _celulas(linha: str) -> list:
    return [c.strip() for c in linha.strip().strip("|").split("|")]


def parse(texto: str) -> Manual:
    meta, corpo = _parse_front_matter(texto)
    linhas = corpo.splitlines()
    blocos: list = []
    buffer: list = []
    i = 0

    def descarregar():
        if buffer:
            blocos.append(Paragrafo(" ".join(l.strip() for l in buffer)))
            buffer.clear()

    while i < len(linhas):
        linha = linhas[i]

        if linha.startswith("```"):
            descarregar()
            linguagem = linha[3:].strip()
            i += 1
            corpo_codigo = []
            while i < len(linhas) and not linhas[i].startswith("```"):
                corpo_codigo.append(linhas[i])
                i += 1
            if i >= len(linhas):
                raise ErroDeFonte("bloco de código não fechado (falta ```)")
            blocos.append(Codigo(linguagem, corpo_codigo))
            i += 1
            continue

        cabecalho = re.match(r"^(#{2,4})\s+(.*)$", linha)
        if cabecalho:
            descarregar()
            blocos.append(Titulo(len(cabecalho.group(1)) - 1, cabecalho.group(2).strip()))
            i += 1
            continue

        if linha.lstrip().startswith("|") and i + 1 < len(linhas) and set(
            linhas[i + 1].replace("|", "").replace(":", "").replace(" ", "").strip()
        ) == {"-"}:
            descarregar()
            cabecalho_tabela = _celulas(linha)
            i += 2
            corpo_tabela = []
            while i < len(linhas) and linhas[i].lstrip().startswith("|"):
                celulas = _celulas(linhas[i])
                if len(celulas) != len(cabecalho_tabela):
                    raise ErroDeFonte(
                        "tabela com número de colunas inconsistente na linha: %s" % linhas[i]
                    )
                corpo_tabela.append(celulas)
                i += 1
            blocos.append(Tabela(cabecalho_tabela, corpo_tabela))
            continue

        if linha.startswith(">"):
            descarregar()
            conteudo = linha[1:].strip()
            achou = re.match(r"^\*\*([^:*]+):\*\*\s*(.*)$", conteudo)
            if achou and achou.group(1) in ROTULOS_CALLOUT:
                blocos.append(Callout(achou.group(1), achou.group(2)))
            else:
                blocos.append(Callout("", conteudo))
            i += 1
            continue

        item = re.match(r"^(-|\d+\.)\s+(.*)$", linha)
        if item:
            descarregar()
            ordenada = item.group(1) != "-"
            itens = []
            while i < len(linhas):
                atual = re.match(r"^(-|\d+\.)\s+(.*)$", linhas[i])
                if not atual or (atual.group(1) != "-") != ordenada:
                    break
                itens.append(atual.group(2).strip())
                i += 1
            blocos.append(Lista(itens, ordenada))
            continue

        if not linha.strip():
            descarregar()
        else:
            buffer.append(linha)
        i += 1

    descarregar()
    return Manual(meta=meta, blocos=blocos)


def numerar(blocos: list) -> list:
    """Preenche Titulo.numero in-place com numeração hierárquica.

    Levanta ErroDeFonte se algum título pular um nível.
    """
    contadores = [0, 0, 0]
    for bloco in blocos:
        if not isinstance(bloco, Titulo):
            continue
        indice = bloco.nivel - 1
        if indice > 0 and contadores[indice - 1] == 0:
            raise ErroDeFonte(
                "título '%s' pula nível (esperado nível %d, encontrado %d)"
                % (bloco.texto, indice, bloco.nivel)
            )
        contadores[indice] += 1
        for maior in range(indice + 1, 3):
            contadores[maior] = 0
        bloco.numero = ".".join(str(c) for c in contadores[: indice + 1])
    return blocos


def _texto_dos_blocos(manual: Manual) -> str:
    """Extrai todo o texto dos blocos de um manual."""
    partes = []
    for bloco in manual.blocos:
        if isinstance(bloco, (Paragrafo, Titulo)):
            partes.append(bloco.texto)
        elif isinstance(bloco, Callout):
            partes.append(bloco.texto)
        elif isinstance(bloco, Lista):
            partes.extend(bloco.itens)
        elif isinstance(bloco, Tabela):
            partes.extend(bloco.cabecalho)
            for linha in bloco.linhas:
                partes.extend(linha)
    return "\n".join(partes)


def validar(manual: Manual, fontes_conhecidas: set) -> list:
    """Valida o manual e retorna lista de avisos.

    Levanta ErroDeFonte se houver referência cruzada quebrada.
    """
    texto = _texto_dos_blocos(manual)
    avisos = []
    if "CONFIRMAR" in texto:
        pendentes = texto.count("CONFIRMAR")
        avisos.append("%d marcação(ões) ⚠️ CONFIRMAR ainda pendentes" % pendentes)
    referencias = set(re.findall(r"\b(\d{2}-[a-z0-9-]+\.md)\b", texto))
    quebradas = sorted(r for r in referencias if r not in fontes_conhecidas)
    if quebradas:
        raise ErroDeFonte("referência para manual inexistente: %s" % ", ".join(quebradas))
    return avisos


def nome_saida(nome_fonte: str) -> str:
    """Deriva o nome do arquivo de saída (.docx) a partir do nome da fonte."""
    base = re.sub(r"^\d+-", "", nome_fonte).removesuffix(".md")
    palavras = "-".join(p.capitalize() for p in base.split("-"))
    return "SCPI-Manual-%s.docx" % palavras


def _sombrear(paragrafo, cor_hex: str) -> None:
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), cor_hex)
    paragrafo._p.get_or_add_pPr().append(sombra)


def _campo_sumario(paragrafo) -> None:
    """Insere o campo TOC do Word (preenche ao 'Atualizar Campo')."""
    for tipo, texto in (("begin", None), (None, r'TOC \o "1-3" \h \z \u'), ("end", None)):
        run = paragrafo.add_run()
        if tipo:
            marca = OxmlElement("w:fldChar")
            marca.set(qn("w:fldCharType"), tipo)
            run._r.append(marca)
        else:
            instrucao = OxmlElement("w:instrText")
            instrucao.set(qn("xml:space"), "preserve")
            instrucao.text = texto
            run._r.append(instrucao)


def _escrever_inline(paragrafo, texto: str) -> None:
    for estilo, pedaco in partir_inline(texto):
        run = paragrafo.add_run(pedaco)
        if estilo == "negrito":
            run.bold = True
        elif estilo == "codigo":
            run.font.name = FONTE_CODIGO
            run.font.size = Pt(10)


def gerar_docx(manual: Manual, destino: Path) -> None:
    documento = docx.Document()

    capa = documento.add_paragraph()
    capa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo = capa.add_run(manual.meta["titulo"] + " — SCPI")
    titulo.bold = True
    titulo.font.size = Pt(24)

    sub = documento.add_paragraph(manual.meta["subtitulo"])
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assinatura = documento.add_paragraph(
        "Versão %s  •  %s  •  Equipe SCPI" % (manual.meta["versao"], manual.meta["data"])
    )
    assinatura.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_page_break()
    documento.add_paragraph("Sumário", style="Heading 1")
    documento.add_paragraph(
        "Após abrir no Word: clique com o botão direito sobre o sumário e escolha "
        "“Atualizar Campo” → “Atualizar o índice inteiro” para preencher as páginas."
    )
    _campo_sumario(documento.add_paragraph())
    documento.add_page_break()

    for bloco in manual.blocos:
        if isinstance(bloco, Titulo):
            rotulo = ("%s %s" % (bloco.numero, bloco.texto)).strip()
            documento.add_paragraph(rotulo, style="Heading %d" % bloco.nivel)
        elif isinstance(bloco, Paragrafo):
            _escrever_inline(documento.add_paragraph(), bloco.texto)
        elif isinstance(bloco, Lista):
            estilo = "List Number" if bloco.ordenada else "List Bullet"
            for item in bloco.itens:
                _escrever_inline(documento.add_paragraph(style=estilo), item)
        elif isinstance(bloco, Callout):
            paragrafo = documento.add_paragraph()
            if bloco.rotulo:
                marca = paragrafo.add_run(bloco.rotulo + ": ")
                marca.bold = True
                marca.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
            _escrever_inline(paragrafo, bloco.texto)
            _sombrear(paragrafo, "FFF6E5")
        elif isinstance(bloco, Codigo):
            for linha in bloco.linhas:
                paragrafo = documento.add_paragraph()
                run = paragrafo.add_run(linha)
                run.font.name = FONTE_CODIGO
                run.font.size = Pt(9)
                _sombrear(paragrafo, CINZA_FUNDO)
        elif isinstance(bloco, Tabela):
            tabela = documento.add_table(rows=1, cols=len(bloco.cabecalho))
            tabela.style = "Table Grid"
            for celula, texto in zip(tabela.rows[0].cells, bloco.cabecalho):
                celula.text = ""
                run = celula.paragraphs[0].add_run(texto)
                run.bold = True
            for linha in bloco.linhas:
                celulas = tabela.add_row().cells
                for celula, texto in zip(celulas, linha):
                    celula.text = ""
                    _escrever_inline(celula.paragraphs[0], texto)
            documento.add_paragraph()

    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(destino))


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    fontes_conhecidas = {p.name for p in sorted(RAIZ_DOCS.glob("[0-9][0-9]-*.md"))}
    alvos = [RAIZ_DOCS / nome for nome in argv] or sorted(
        RAIZ_DOCS.glob("[0-9][0-9]-*.md")
    )
    if not alvos:
        print("nenhuma fonte encontrada em %s" % RAIZ_DOCS)
        return 1

    for caminho in alvos:
        try:
            manual = parse(caminho.read_text(encoding="utf-8"))
            avisos = validar(manual, fontes_conhecidas)
            numerar(manual.blocos)
            saida = DESTINO / nome_saida(caminho.name)
            gerar_docx(manual, saida)
        except ErroDeFonte as erro:
            print("ERRO em %s: %s" % (caminho.name, erro))
            return 1
        for aviso in avisos:
            print("aviso  %s: %s" % (caminho.name, aviso))
        print("gerado %s" % saida.relative_to(RAIZ_DOCS.parent))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
